"""Arma dos planillas de Excel REALES (las escriben bibliotecas de Python, no las armamos a mano) para probar lib/xlsx.ts
y la carga del arancel de punta a punta (e2e/plan-y-arancel.spec.ts):

    python3 lib/__fixtures__/generar-arancel-prueba.py

arancel-prueba.xlsx (openpyxl: guarda el texto «en línea», como hacen muchos programas que exportan a Excel)
- «Hoja1» vacía y primera (la carga abre la primera hoja CON datos).
- «Aranceles»: encabezado «Prestación | Precio», tildes, una fila vacía (la 4), precios numéricos y uno escrito como texto
  («Gs. 450.000»), una fórmula (openpyxl no guarda su resultado: el lector la devuelve como «=B2*10»), un nombre con
  formato mixto (texto enriquecido) y dos prestaciones que ya existen en la demo con otro precio.
- «Notas»: oculta (no se ofrece).
- «Laboratorio»: otra hoja visible, con código, para elegir hoja.

arancel-cadenas-compartidas.xlsx (xlsxwriter: guarda el texto en la tabla de cadenas compartidas, como Excel)
- «Precios»: la tabla empieza en B2 (la columna A vacía no cuenta), un nombre con formato mixto, una fórmula CON su
  resultado guardado, un precio con decimales, un booleano en una columna de más y un nombre repetido.
- «Oculta»: muy oculta (veryHidden).

Si lo cambiás, actualizá lib/xlsx.test.ts y e2e/plan-y-arancel.spec.ts.
"""
from pathlib import Path

import xlsxwriter
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

aca = Path(__file__).parent

# ── openpyxl ──
libro = Workbook()
libro.active.title = "Hoja1"

aranceles = libro.create_sheet("Aranceles")
aranceles.append(["Prestación", "Precio"])
aranceles.append(["Limpieza con ultrasonido", 180000])
aranceles.append(["Extracción de muela del juicio", 950000])
aranceles.append([])
aranceles.append(["Blanqueamiento con férula", "Gs. 450.000"])
aranceles.append(["Carillas de porcelana (por pieza)", "=B2*10"])
aranceles.append(["Profilaxis (adulto)", 270000])
aranceles.append([CellRichText(["Corona ", TextBlock(InlineFont(b=True), "de porcelana"), "/cerámica"]), 2900000])
aranceles.append(["Ortodoncia — control mensual", 350000])

notas = libro.create_sheet("Notas")
notas.append(["No cargar", 1])
notas.sheet_state = "hidden"

laboratorio = libro.create_sheet("Laboratorio")
laboratorio.append(["Código", "Descripción", "Precio"])
laboratorio.append(["LAB01", "Corona de zirconio (laboratorio)", 1200000])

libro.save(aca / "arancel-prueba.xlsx")

# ── xlsxwriter ──
libro = xlsxwriter.Workbook(aca / "arancel-cadenas-compartidas.xlsx")
precios = libro.add_worksheet("Precios")
negrita = libro.add_format({"bold": True})
precios.write("B2", "Prestación")
precios.write("C2", "Precio")
precios.write("B3", "Limpieza con ultrasonido")
precios.write_number("C3", 180000)
precios.write_rich_string("B4", "Corona ", negrita, "de porcelana", "/cerámica")
precios.write_formula("C4", "=C3*16", None, 2880000)
precios.write("B5", "Extracción de muela del juicio")
precios.write_number("C5", 950000.5)
precios.write("B6", "Ortodoncia — control mensual")
precios.write("C6", "Gs. 350.000")
precios.write_boolean("D6", True)
precios.write("B7", "Limpieza con ultrasonido")
precios.write_number("C7", 190000)
oculta = libro.add_worksheet("Oculta")
oculta.write("A1", "No cargar")
oculta.very_hidden()
libro.close()
